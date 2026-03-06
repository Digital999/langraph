"""报告生成 Agent - 生成 Word 报告"""
import json
import os
import traceback
from datetime import datetime
from typing import Any, Dict

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from langchain_core.prompts import ChatPromptTemplate
from langchain_openai import ChatOpenAI
from langgraph.config import get_stream_writer

from config import settings
from models.schemas import AgentState
from utils.constants import LLM_MAX_RETRIES, LLM_TIMEOUT, REPORT_LLM_TEMPERATURE
from utils.logger import logger

REPORT_SYSTEM_PROMPT = """你是一个专业的报告生成助手，负责将查询结果整理成结构化的报告内容。

你的职责：
1. 分析查询结果，提取关键信息
2. 组织报告结构，包括：标题、摘要、详细信息、结论
3. 使用清晰、专业的语言描述数据
4. 对敏感信息进行适当处理
5. 生成易读的报告文本

报告结构要求：
- 报告标题：简洁明了
- 查询摘要：概述查询目的和范围
- 详细信息：分类展示查询结果
- 数据分析：对结果进行简要分析
- 结论建议：给出相关建议（如适用）

输出格式：
返回JSON格式：{{"title": "标题", "summary": "摘要", "details": "详细内容", "conclusion": "结论"}}
"""

llm = ChatOpenAI(
    model=settings.MODEL_NAME,
    api_key=settings.OPENAI_API_KEY,
    base_url=settings.OPENAI_BASE_URL,
    temperature=REPORT_LLM_TEMPERATURE,
    max_retries=LLM_MAX_RETRIES,
    timeout=LLM_TIMEOUT
)

prompt = ChatPromptTemplate.from_messages([
    ("system", REPORT_SYSTEM_PROMPT),
    ("human", "用户查询：{user_input}\n\n查询结果：{query_results}\n\n请生成报告内容。")
])

chain = prompt | llm

def generate_word_report(report_content: Dict[str, Any], filename: str) -> str:
    """
    生成Word报告文件
    
    Args:
        report_content: 报告内容字典
        filename: 文件名
        
    Returns:
        文件路径
    """
    doc = Document()
    
    # 标题
    title = doc.add_heading(report_content.get("title", "用户信息查询报告"), 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    # 生成时间
    doc.add_paragraph(f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph("")
    
    # 摘要
    doc.add_heading("查询摘要", 1)
    doc.add_paragraph(report_content.get("summary", ""))
    doc.add_paragraph("")
    
    # 详细信息
    doc.add_heading("详细信息", 1)
    details = report_content.get("details", "")
    
    # 如果 details 是字典，格式化输出
    if isinstance(details, dict):
        for section, content in details.items():
            doc.add_heading(section, 2)
            if isinstance(content, dict):
                for key, value in content.items():
                    doc.add_paragraph(f"{key}: {value}", style='List Bullet')
            elif isinstance(content, list):
                for item in content:
                    doc.add_paragraph(str(item), style='List Bullet')
            else:
                doc.add_paragraph(str(content))
            doc.add_paragraph("")
    else:
        doc.add_paragraph(str(details))
        doc.add_paragraph("")
    
    # 结论
    doc.add_heading("结论", 1)
    conclusion = report_content.get("conclusion", "")
    if isinstance(conclusion, dict):
        for key, value in conclusion.items():
            doc.add_paragraph(f"{key}: {value}")
    else:
        doc.add_paragraph(str(conclusion))
    
    # 保存文件
    os.makedirs("reports", exist_ok=True)
    filepath = f"reports/{filename}"
    doc.save(filepath)
    return filepath

def process_report(state: AgentState) -> AgentState:
    """生成报告"""
    logger.separator()
    logger.info("开始生成报告")
    
    # 获取流式写入器
    try:
        writer = get_stream_writer()
    except Exception:
        writer = None
    
    if not state.get("query_results"):
        logger.error("缺少查询结果")
        state["error"] = "缺少查询结果"
        state["next_step"] = "end"
        if writer:
            writer({"type": "content", "content": "\n\n缺少查询结果"})
        return state
    
    logger.debug(f"查询结果: {state['query_results']}")
    
    try:
        # 推送状态：开始生成报告
        if writer:
            writer({"type": "content", "content": "正在生成报告..."})
        
        # 使用 LLM 生成报告内容
        logger.debug("调用 LLM 生成报告内容...")
        start_time = datetime.now()
        
        response = chain.invoke({
            "user_input": state.get("user_input", ""),
            "query_results": str(state.get("query_results", {}))
        })
        
        elapsed = (datetime.now() - start_time).total_seconds() * 1000
        logger.debug(f"LLM 响应耗时: {elapsed:.0f}ms")
        
        # 处理响应内容
        content: str = response.content if isinstance(response.content, str) else str(response.content)
        if isinstance(response.content, list) and response.content:
            content = str(response.content[0])
        
        content_str = content.strip()
        logger.llm_output(content_str)
        
        # 清理 markdown 代码块标记
        if content_str.startswith("```json"):
            content_str = content_str[7:]
        elif content_str.startswith("```"):
            content_str = content_str[3:]
        
        if content_str.endswith("```"):
            content_str = content_str[:-3]
        
        content_str = content_str.strip()
        
        report_content = json.loads(content_str)
        logger.debug(f"报告结构: {list(report_content.keys())}")
        
        # 推送状态：正在生成Word文件
        if writer:
            writer({"type": "status", "message": "正在生成Word文件..."})
        
        # 生成 Word 文件
        filename = f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        logger.info(f"生成 Word 文件: {filename}")
        
        filepath = generate_word_report(report_content, filename)
        
        state["report_path"] = filepath
        state["next_step"] = "end"
        logger.info(f"✓ 报告生成成功: {filepath}")
        
        # 推送成功消息和下载链接
        if writer:
            # 先推送成功消息（作为普通内容）
            writer({"type": "message", "content": "\n✓ 报告生成成功！"})
            # 再推送报告URL（前端会渲染为下载按钮）
            report_url = f"/api/download/{filename}"
            writer({"type": "report", "filename": filename, "url": report_url})
    
    except Exception:
        logger.error(f"报告生成失败:\n{traceback.format_exc()}")
        state["error"] = "报告生成失败，请稍后重试"
        state["next_step"] = "end"
        if writer:
            writer({"type": "content", "content": "\n\n" + state["error"]})
    
    return state
